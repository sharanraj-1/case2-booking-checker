"""Build test_bookings.xlsx: the judges' template layout with one planted mistake per rule.

Layout: row 1 title, row 2 headers, row 3 grey example (BK-001), data from row 4.
Run: python make_test_file.py
"""

from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill

BOOKING_HEADERS = [
    "Booking ID", "Seller name", "Origin city", "Origin country", "Destination city",
    "Destination country", "Destination type (Warehouse / Business address)",
    "Ready to ship date (DD-MM-YYYY)", "Seller priority (Cheapest / Fastest / Most reliable / Balanced)",
    "Pallets", "Cartons", "Total declared value", "Currency", "Total gross weight (kg)",
    "Delivery appointment booked (Yes / No)",
]
ITEM_HEADERS = [
    "Booking ID", "SKU", "Description", "HS code", "Quantity (units)", "Unit price",
    "Units per carton", "Cartons", "Unit weight (kg)", "Country of origin",
]


def item(sku, desc, hs, qty, price, upc, cartons, unit_kg, coo):
    return dict(sku=sku, desc=desc, hs=hs, qty=qty, price=price, upc=upc, cartons=cartons, kg=unit_kg, coo=coo)


def booking(bid, seller, origin, city, country, dtype, ready, prio, pallets, items, appt="Yes",
            gross=None, declared=None, cartons=None, packaging_per_pallet=25):
    """A clean booking whose totals are derived from its items, unless overridden."""
    net = sum(i["qty"] * i["kg"] for i in items)
    return dict(
        bid=bid, seller=seller, origin=origin, city=city, country=country, dtype=dtype, ready=ready, prio=prio,
        pallets=pallets, items=items, appt=appt,
        cartons=cartons if cartons is not None else sum(i["cartons"] for i in items),
        declared=declared if declared is not None else round(sum(i["qty"] * i["price"] for i in items), 2),
        gross=gross if gross is not None else round(net + packaging_per_pallet * (pallets or 1)),
    )


CASES = [
    # Example row from the template: passes all 10 rules; Cheapest -> Carrier A at £540.
    booking("BK-001", "Northfield Goods", "London", "Berlin", "Germany", "Warehouse", "05-11-2026", "Cheapest", 3,
            [item("NG-MUG-01", "Ceramic coffee mug", 691200, 1080, 5, 12, 90, 0.5, "United Kingdom")], gross=620),
    # R-01: blank HS, 5-digit HS (lost leading zero) and 4-digit HS.
    booking("BK-002", "Spice & Co", "Manchester", "Paris", "France", "Business address", datetime(2026, 11, 6), "Cheapest", 1,
            [item("SC-PEP", "Black peppercorns 250g", 90412, 400, 3.5, 20, 20, 0.25, "India"),
             item("SC-TEE", "Cotton t-shirt", None, 200, 8, 25, 8, 0.2, "Turkey"),
             item("SC-SCF", "Wool scarf", 6117, 100, 12, 50, 2, 0.15, "China")], appt="No"),
    # R-02: declared value 5% below the items total.
    booking("BK-003", "Brightside Homeware", "Birmingham", "Hamburg", "Germany", "Warehouse", "09-11-2026", "Cheapest", 2,
            [item("BH-LAMP", "Ceramic table lamp", 940520, 240, 25, 6, 40, 1.5, "China")], declared=5700),
    # R-02 tolerance: 0.8% gap passes. Also checks "battery-free" and candles are NOT flagged.
    booking("BK-004", "Leeds Living", "Leeds", "Amsterdam", "Netherlands", "Business address", "10-11-2026", "Cheapest", 1,
            [item("LL-TIMER", "Battery-free kitchen timer", 910700, 300, 6, 30, 10, 0.1, "China"),
             item("LL-CNDL", "Scented soy candle", 340600, 200, 9, 20, 10, 0.3, "United Kingdom")],
            appt="No", declared=3628.80),
    # R-03: 12 units/carton x 20 cartons = 240, but quantity is 250.
    booking("BK-005", "Pennine Textiles", "Leeds", "Frankfurt", "Germany", "Warehouse", "12-11-2026", "Fastest", 1,
            [item("PT-TOWEL", "Cotton bath towel", 630260, 250, 7, 12, 20, 0.6, "India")]),
    # R-04: booking says 100 cartons, items add up to 90.
    booking("BK-006", "Mersey Glassworks", "Manchester", "Rotterdam", "Netherlands", "Warehouse", "12-11-2026", "Most reliable", 2,
            [item("MG-JUG", "Glass water jug", 701349, 360, 10, 4, 90, 1.2, "Turkey")], cartons=100),
    # R-05: 2,400 kg on 3 pallets = 800 kg/pallet; no carrier fits. B needs 4 (£880), A 5 (£900), C 6 (£2,040).
    booking("BK-007", "Heavy Stone Ltd", "Birmingham", "Berlin", "Germany", "Warehouse", "13-11-2026", "Cheapest", 3,
            [item("HS-TILE", "Porcelain floor tiles", 690721, 1000, 4.5, 10, 100, 2.3, "India")], gross=2400),
    # R-05 fallback: 450 kg/pallet excludes Carrier C, so Fastest falls back to Carrier B.
    booking("BK-008", "Oakline Furniture", "London", "Lyon", "France", "Business address", "16-11-2026", "Fastest", 2,
            [item("OF-CHAIR", "Oak dining chair", 940161, 40, 85, 2, 20, 21, "United Kingdom")], appt="No", gross=900),
    # R-06: power bank and perfume.
    booking("BK-009", "Gadget Hub", "London", "Berlin", "Germany", "Business address", "16-11-2026", "Cheapest", 1,
            [item("GH-PB", "Power bank 10000mAh", 850760, 100, 15, 20, 5, 0.25, "China"),
             item("GH-PERF", "Perfume 50ml gift set", 330300, 60, 20, 12, 5, 0.3, "Vietnam"),
             item("GH-CASE", "Silicone phone case", 392690, 200, 4, 50, 4, 0.05, "China")], appt="No"),
    # R-07: gross 230 kg but goods alone weigh 270 kg.
    booking("BK-010", "Northern Books", "Manchester", "Paris", "France", "Business address", "17-11-2026", "Most reliable", 1,
            [item("NB-BOOK", "Hardback cookbook", 490199, 300, 12, 15, 20, 0.9, "United Kingdom")], appt="No", gross=230),
    # R-08: warehouse delivery with no appointment. HS stored as text keeps its leading zero.
    booking("BK-011", "Thames Tea Co", "London", "Amsterdam", "Netherlands", "Warehouse", "17-11-2026", "Balanced", 1,
            [item("TT-TEA", "Loose leaf green tea 100g", "090210", 600, 4, 24, 25, 0.12, "India")], appt="No"),
    # R-09: blank country of origin.
    booking("BK-012", "Midlands Toys", "Birmingham", "Paris", "France", "Warehouse", "18-11-2026", "Cheapest", 1,
            [item("MT-PUZ", "Wooden jigsaw puzzle", 950490, 150, 9, 15, 10, 0.6, None)]),
    # R-10: ready date in the past.
    booking("BK-013", "Yorkshire Knitwear", "Leeds", "Hamburg", "Germany", "Business address", "15-09-2026", "Balanced", 1,
            [item("YK-JUMP", "Wool jumper", 611011, 120, 30, 12, 10, 0.5, "United Kingdom")], appt="No"),
    # Balanced with all three carriers allowed -> Carrier B (minimax regret).
    booking("BK-014", "Cotswold Ceramics", "London", "Lyon", "France", "Warehouse", datetime(2026, 11, 19), "Balanced", 2,
            [item("CC-BOWL", "Stoneware bowl", 691200, 480, 6, 12, 40, 0.7, "United Kingdom")]),
    # Most reliable, light load -> Carrier C.
    booking("BK-015", "Harbour Prints", "Manchester", "Rotterdam", "Netherlands", "Business address", "19-11-2026", "Most reliable", 1,
            [item("HP-PRINT", "Framed art print", 491191, 50, 40, 5, 10, 1.8, "United Kingdom")], appt="No"),
    # Ambiguous R-06 word: advisory only, booking stays Ready.
    booking("BK-016", "Desk Supplies Co", "Birmingham", "Frankfurt", "Germany", "Warehouse", "20-11-2026", "Cheapest", 1,
            [item("DS-MOUSE", "Wireless mouse", 847160, 200, 11, 20, 10, 0.1, "China")]),
    # Trap: pallets left blank -> R-05 Red with pallets needed per carrier.
    booking("BK-017", "Lakeside Linens", "Leeds", "Berlin", "Germany", "Business address", "20-11-2026", "Cheapest", None,
            [item("LL-SHEET", "Linen bed sheet", 630231, 120, 22, 10, 12, 0.9, "Turkey")], appt="No"),
    # Trap: booking with no item lines.
    booking("BK-018", "Ghost Orders Ltd", "London", "Paris", "France", "Business address", "20-11-2026", "Cheapest", 1,
            [], appt="No", gross=100, declared=1000, cartons=10),
]

# An item whose booking isn't on the Bookings sheet.
ORPHAN_ITEMS = [dict(bid="BK-099", sku="ZZ-ORPHAN", desc="Canvas tote bag", hs=420222, qty=50, price=3,
                     upc=25, cartons=2, kg=0.2, coo="India")]


def write(path="test_bookings.xlsx"):
    wb = openpyxl.Workbook()
    grey = PatternFill("solid", fgColor="D9D9D9")
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Bookings"
    ws.append(["Bookings – UK to EU pallet shipments (one row per booking)"])
    ws.append(BOOKING_HEADERS)
    for c in ws[2]:
        c.font = bold
    for n, b in enumerate(CASES):
        ws.append([b["bid"], b["seller"], b["origin"], "United Kingdom", b["city"], b["country"], b["dtype"],
                   b["ready"], b["prio"], b["pallets"], b["cartons"], b["declared"], "GBP", b["gross"], b["appt"]])
        if n == 0:
            for c in ws[ws.max_row]:
                c.fill = grey
        if n == 4:
            ws.append([None] * len(BOOKING_HEADERS))  # stray blank row

    wi = wb.create_sheet("Items")
    wi.append(["Items – one row per SKU per booking"])
    wi.append(ITEM_HEADERS)
    for c in wi[2]:
        c.font = bold
    first = True
    for b in CASES:
        for i in b["items"]:
            wi.append([b["bid"], i["sku"], i["desc"], i["hs"], i["qty"], i["price"], i["upc"], i["cartons"], i["kg"], i["coo"]])
            if first:
                for c in wi[wi.max_row]:
                    c.fill = grey
                first = False
    for i in ORPHAN_ITEMS:
        wi.append([i["bid"], i["sku"], i["desc"], i["hs"], i["qty"], i["price"], i["upc"], i["cartons"], i["kg"], i["coo"]])
    for row in wi.iter_rows(min_row=3, min_col=4, max_col=4):
        for c in row:
            if isinstance(c.value, str):
                c.number_format = "@"
    for sheet in (ws, wi):
        for col in sheet.columns:
            sheet.column_dimensions[col[1].column_letter].width = 18
    wb.save(path)
    return path


if __name__ == "__main__":
    print("Wrote", write())
